"""Read-only document extraction helpers for Archive Search V11."""

from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path
import os
import shutil
import stat
import tempfile
from typing import Dict, Iterator, List, Optional, Sequence

import openpyxl

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    xlrd = None
    HAS_XLRD = False

try:
    from pyxlsb import open_workbook as open_xlsb_workbook
    HAS_PYXLSB = True
except ImportError:
    open_xlsb_workbook = None
    HAS_PYXLSB = False

from archive_config import OPENPYXL_EXTENSIONS, PLAIN_TEXT_EXTENSIONS, XLSB_EXTENSIONS
from archive_file_safety import detect_supported_file_type, ensure_binary_readonly_access
from archive_optional import DocxDocument, HAS_DOCX, HAS_PDF, HAS_WIN32_COM, PdfReader, pythoncom, win32com
from archive_utils import compact_whitespace, display_text


def convert_xls_value(value, ctype, datemode):
    """Convert an xlrd cell value into a more useful Python value."""
    if not HAS_XLRD:
        return value
    if ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(value, datemode)
        except Exception:
            return value
    if ctype == xlrd.XL_CELL_NUMBER:
        if float(value).is_integer():
            return int(value)
        return value
    if ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(value)
    if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    return value


@contextmanager
def temporary_word_source_copy(file_path: Path) -> Iterator[Path]:
    """Copy a Word source file to a temporary path and yield that copy."""
    ensure_binary_readonly_access(file_path)
    temp_dir = Path(tempfile.mkdtemp(prefix="archive_search_word_"))
    temp_path = temp_dir / f"source{file_path.suffix.lower() or '.doc'}"
    try:
        with file_path.open("rb") as source_handle, temp_path.open("wb") as temp_handle:
            shutil.copyfileobj(source_handle, temp_handle, length=1024 * 1024)
        try:
            os.chmod(temp_path, stat.S_IREAD)
        except Exception:
            pass
        yield temp_path
    finally:
        try:
            os.chmod(temp_path, stat.S_IWRITE | stat.S_IREAD)
        except Exception:
            pass
        try:
            temp_path.unlink(missing_ok=True)
        except Exception:
            pass
        try:
            temp_dir.rmdir()
        except Exception:
            pass


class ReadOnlyWordSession:
    """Shared read-only Microsoft Word session for legacy .doc extraction."""

    def __init__(self) -> None:
        self.word = None

    def __enter__(self) -> "ReadOnlyWordSession":
        if not HAS_WIN32_COM:
            raise RuntimeError("pywin32 is not installed.")
        pythoncom.CoInitialize()
        self.word = win32com.client.DispatchEx("Word.Application")
        self.word.Visible = False
        self.word.DisplayAlerts = 0
        try:
            self.word.AutomationSecurity = 3
        except Exception:
            pass
        try:
            self.word.ScreenUpdating = False
        except Exception:
            pass
        return self

    def __exit__(self, exc_type, exc, tb) -> None:
        if self.word is not None:
            try:
                self.word.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass

    def extract_lines(self, file_path: Path) -> List[str]:
        if self.word is None:
            raise RuntimeError("Word session is not open.")
        document = None
        with temporary_word_source_copy(file_path) as safe_input_path:
            try:
                document = self.word.Documents.Open(
                    str(safe_input_path.resolve()),
                    ConfirmConversions=False,
                    ReadOnly=True,
                    AddToRecentFiles=False,
                    Visible=False,
                    Revert=False,
                    OpenAndRepair=False,
                    NoEncodingDialog=True,
                )
                raw_text = document.Content.Text or ""
                raw_text = raw_text.replace("\r", "\n").replace("\x07", " ")
                lines = [compact_whitespace(line) for line in raw_text.splitlines()]
                return [line for line in lines if line]
            finally:
                if document is not None:
                    try:
                        document.Close(False)
                    except Exception:
                        pass


class DocumentExtractor:
    """Read-only extraction for every supported file type."""

    @staticmethod
    def infer_file_type(file_path: Path) -> str:
        detected = detect_supported_file_type(file_path)
        return detected or "unknown"

    @staticmethod
    def build_row_text(row_values: Sequence[object]) -> Optional[str]:
        display_cells = [display_text(value) for value in row_values]
        while display_cells and not display_cells[-1]:
            display_cells.pop()
        if not any(display_cells):
            return None
        # Preserve leading and internal blank cells so future indexes keep the
        # worksheet column positions needed for clearer result display.
        return " | ".join(display_cells)

    @staticmethod
    def _clean_lines(text: str) -> List[str]:
        return [line for line in (compact_whitespace(part) for part in text.splitlines()) if line]

    @staticmethod
    def _chunk_paragraphs(lines: Sequence[str], min_join_length: int = 120) -> List[str]:
        """Join short adjacent lines into more readable search chunks."""
        chunks: List[str] = []
        buffer: List[str] = []
        buffer_length = 0
        for line in lines:
            buffer.append(line)
            buffer_length += len(line) + 1
            if buffer_length >= min_join_length or line.endswith((".", ":", ";", "?", "!")):
                chunks.append(" ".join(buffer).strip())
                buffer = []
                buffer_length = 0
        if buffer:
            chunks.append(" ".join(buffer).strip())
        return [chunk for chunk in chunks if chunk]

    def extract_xlsx_chunks(self, file_path: Path) -> List[Dict[str, object]]:
        ensure_binary_readonly_access(file_path)
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
        try:
            chunks: List[Dict[str, object]] = []
            for ws in wb.worksheets:
                for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    row_text = self.build_row_text(row)
                    if not row_text:
                        continue
                    chunks.append(
                        {
                            "location_label": f"{ws.title} | Row {row_number}",
                            "location_sort_key": row_number,
                            "sheet_name": ws.title,
                            "row_number": row_number,
                            "line_number": None,
                            "page_number": None,
                            "content": row_text,
                        }
                    )
            return chunks
        finally:
            wb.close()

    def extract_xls_chunks(self, file_path: Path) -> List[Dict[str, object]]:
        if not HAS_XLRD:
            raise RuntimeError("xlrd is not installed, so legacy .xls files cannot be indexed.")
        ensure_binary_readonly_access(file_path)
        book = xlrd.open_workbook(file_path, on_demand=True)
        try:
            datemode = book.datemode
            chunks: List[Dict[str, object]] = []
            for sheet_name in book.sheet_names():
                sheet = book.sheet_by_name(sheet_name)
                for row_idx in range(sheet.nrows):
                    raw_values = sheet.row_values(row_idx)
                    raw_types = sheet.row_types(row_idx)
                    row_values = [
                        convert_xls_value(value, ctype, datemode)
                        for value, ctype in zip(raw_values, raw_types)
                    ]
                    row_text = self.build_row_text(row_values)
                    if not row_text:
                        continue
                    chunks.append(
                        {
                            "location_label": f"{sheet_name} | Row {row_idx + 1}",
                            "location_sort_key": row_idx + 1,
                            "sheet_name": sheet_name,
                            "row_number": row_idx + 1,
                            "line_number": None,
                            "page_number": None,
                            "content": row_text,
                        }
                    )
            return chunks
        finally:
            try:
                book.release_resources()
            except Exception:
                pass

    def extract_xlsb_chunks(self, file_path: Path) -> List[Dict[str, object]]:
        if not HAS_PYXLSB:
            raise RuntimeError("pyxlsb is not installed, so .xlsb files cannot be indexed.")
        ensure_binary_readonly_access(file_path)
        chunks: List[Dict[str, object]] = []
        with open_xlsb_workbook(str(file_path)) as workbook:
            for sheet_name in workbook.sheets:
                with workbook.get_sheet(sheet_name) as sheet:
                    for row in sheet.rows():
                        row_number = int(getattr(row[0], 'r', 0) or 0) if row else 0
                        row_values = [getattr(cell, 'v', None) for cell in row]
                        row_text = self.build_row_text(row_values)
                        if not row_text:
                            continue
                        chunks.append(
                            {
                                "location_label": f"{sheet_name} | Row {row_number or '?'}",
                                "location_sort_key": row_number or 0,
                                "sheet_name": sheet_name,
                                "row_number": row_number or None,
                                "line_number": None,
                                "page_number": None,
                                "content": row_text,
                            }
                        )
        return chunks

    def extract_docx_paragraphs(self, file_path: Path) -> List[str]:
        if not HAS_DOCX:
            raise RuntimeError("python-docx is not installed.")
        ensure_binary_readonly_access(file_path)
        document = DocxDocument(str(file_path))
        paragraphs: List[str] = []
        for paragraph in document.paragraphs:
            text = compact_whitespace(paragraph.text)
            if text:
                paragraphs.append(text)
        for table in document.tables:
            for row in table.rows:
                cell_values = [compact_whitespace(cell.text) for cell in row.cells]
                row_text = " | ".join(value for value in cell_values if value)
                if row_text:
                    paragraphs.append(row_text)
        return paragraphs

    def extract_word_lines(self, file_path: Path, word_session: Optional[ReadOnlyWordSession] = None) -> List[str]:
        detected_type = detect_supported_file_type(file_path)
        if detected_type != "word":
            raise RuntimeError(f"File signature does not match a supported Word document. File: {file_path.name}; folder: {file_path.parent}; full path: {file_path}")

        if file_path.suffix.lower() == ".doc":
            if word_session is not None:
                return word_session.extract_lines(file_path)
            if not HAS_WIN32_COM:
                raise RuntimeError(".doc files require pywin32 and Microsoft Word.")
            with ReadOnlyWordSession() as session:
                return session.extract_lines(file_path)

        if HAS_DOCX:
            return self.extract_docx_paragraphs(file_path)

        if HAS_WIN32_COM:
            if word_session is not None:
                return word_session.extract_lines(file_path)
            with ReadOnlyWordSession() as session:
                return session.extract_lines(file_path)

        raise RuntimeError(".docx files require python-docx or pywin32 + Microsoft Word.")

    def extract_word_chunks(self, file_path: Path, word_session: Optional[ReadOnlyWordSession] = None) -> List[Dict[str, object]]:
        lines = self.extract_word_lines(file_path, word_session=word_session)
        chunks: List[Dict[str, object]] = []
        for line_number, paragraph in enumerate(lines, start=1):
            chunks.append(
                {
                    "location_label": f"Line {line_number}",
                    "location_sort_key": line_number,
                    "sheet_name": None,
                    "row_number": None,
                    "line_number": line_number,
                    "page_number": None,
                    "content": paragraph,
                }
            )
        return chunks

    def extract_pdf_chunks(self, file_path: Path) -> List[Dict[str, object]]:
        if not HAS_PDF:
            raise RuntimeError("pypdf is not installed.")
        ensure_binary_readonly_access(file_path)
        reader = PdfReader(str(file_path))
        chunks: List[Dict[str, object]] = []
        for page_index, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            page_lines = self._clean_lines(page_text)
            if not page_lines:
                continue
            for block_number, block in enumerate(self._chunk_paragraphs(page_lines), start=1):
                chunks.append(
                    {
                        "location_label": f"Page {page_index}" if block_number == 1 else f"Page {page_index} | Block {block_number}",
                        "location_sort_key": (page_index * 100000) + block_number,
                        "sheet_name": None,
                        "row_number": None,
                        "line_number": None,
                        "page_number": page_index,
                        "content": block,
                    }
                )
        return chunks

    def extract_text_chunks(self, file_path: Path) -> List[Dict[str, object]]:
        ensure_binary_readonly_access(file_path)
        raw_bytes = file_path.read_bytes()
        decoded_text = None
        encodings = ("utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "cp1252", "latin-1")
        for encoding in encodings:
            try:
                decoded_text = raw_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded_text is None:
            decoded_text = raw_bytes.decode("utf-8", errors="replace")

        lines = self._clean_lines(decoded_text)
        chunks: List[Dict[str, object]] = []
        for line_number, block in enumerate(self._chunk_paragraphs(lines), start=1):
            chunks.append(
                {
                    "location_label": f"Line {line_number}",
                    "location_sort_key": line_number,
                    "sheet_name": None,
                    "row_number": None,
                    "line_number": line_number,
                    "page_number": None,
                    "content": block,
                }
            )
        return chunks

    def build_error_record(self, file_path: Path, size, mtime_ns, message: str) -> Dict[str, object]:
        try:
            file_type = self.infer_file_type(file_path)
        except Exception:
            file_type = "unknown"
        return {
            "size": size,
            "mtime_ns": mtime_ns,
            "suffix": file_path.suffix.lower(),
            "file_type": file_type or "unknown",
            "chunks": [],
            "error": message,
        }

    def build_index_record(
        self,
        file_path: Path,
        size: Optional[int],
        mtime_ns: Optional[int],
        word_session: Optional[ReadOnlyWordSession] = None,
    ) -> Dict[str, object]:
        try:
            detected_type = detect_supported_file_type(file_path)
        except Exception as exc:
            return self.build_error_record(file_path, size, mtime_ns, f"Could not confirm read-only access for {file_path}: {exc}")

        try:
            suffix = file_path.suffix.lower()
            if detected_type == "excel":
                if suffix in OPENPYXL_EXTENSIONS:
                    chunks = self.extract_xlsx_chunks(file_path)
                elif suffix in XLSB_EXTENSIONS:
                    chunks = self.extract_xlsb_chunks(file_path)
                else:
                    chunks = self.extract_xls_chunks(file_path)
                return {
                    "size": size,
                    "mtime_ns": mtime_ns,
                    "suffix": suffix,
                    "file_type": "excel",
                    "chunks": chunks,
                    "error": None,
                }
            if detected_type == "word":
                return {
                    "size": size,
                    "mtime_ns": mtime_ns,
                    "suffix": suffix,
                    "file_type": "word",
                    "chunks": self.extract_word_chunks(file_path, word_session=word_session),
                    "error": None,
                }
            if detected_type == "pdf":
                return {
                    "size": size,
                    "mtime_ns": mtime_ns,
                    "suffix": suffix,
                    "file_type": "pdf",
                    "chunks": self.extract_pdf_chunks(file_path),
                    "error": None,
                }
            if detected_type == "text" or suffix in PLAIN_TEXT_EXTENSIONS:
                return {
                    "size": size,
                    "mtime_ns": mtime_ns,
                    "suffix": suffix,
                    "file_type": "text",
                    "chunks": self.extract_text_chunks(file_path),
                    "error": None,
                }
            return self.build_error_record(
                file_path,
                size,
                mtime_ns,
                f"Unsupported or mismatched file type. File: {file_path.name}; folder: {file_path.parent}; full path: {file_path}",
            )
        except Exception as exc:
            return self.build_error_record(
                file_path,
                size,
                mtime_ns,
                f"Could not open file. File: {file_path.name}; folder: {file_path.parent}; full path: {file_path}; error: {exc}",
            )
